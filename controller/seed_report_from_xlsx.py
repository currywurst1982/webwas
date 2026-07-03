#!/usr/bin/env python3
"""
One-time import: load a "주간 운영 현황" report .xlsx (see attached template)
into the report_store SQLite database.

Usage:
    python3 seed_report_from_xlsx.py /path/to/report.xlsx
"""
import re
import sys

import openpyxl

import report_store


def _s(v) -> str:
    """Cell value -> trimmed string. Internal newlines (e.g. multi-paragraph
    incident write-ups) are preserved so the UI can render them with
    white-space: pre-wrap."""
    if v is None:
        return ""
    return str(v).strip()


def _s_date(v) -> str:
    """Date / date-range cell -> single line, e.g.
    "06월 29일\\n~\\n07월 03일" -> "06월 29일~07월 03일"."""
    if v is None:
        return ""
    text = str(v).strip()
    text = re.sub(r"\s*\n\s*~\s*\n\s*", "~", text)
    text = re.sub(r"\s*\n\s*", " ", text)
    return text


def _n(v) -> int:
    try:
        return int(v)
    except (TypeError, ValueError):
        return 0


def _parse_week_dates(title: str):
    """'금주 진행 내역 (2026.06.30~2026.07.03)' -> ('2026-06-30', '2026-07-03')"""
    m = re.search(r"\((\d{4})\.(\d{2})\.(\d{2})\s*~\s*(\d{4})\.(\d{2})\.(\d{2})\)", title)
    if not m:
        return None, None
    y1, m1, d1, y2, m2, d2 = m.groups()
    return f"{y1}-{m1}-{d1}", f"{y2}-{m2}-{d2}"


def seed(path: str) -> int:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Sheet1"]

    start_date, end_date = _parse_week_dates(str(ws["A9"].value or ""))
    next_start, next_end = _parse_week_dates(str(ws["A26"].value or ""))
    if not start_date:
        print("[ERROR] 금주 진행 내역 제목에서 날짜를 찾을 수 없습니다 (A9 셀 확인).")
        return 1

    week = report_store.create_week(start_date, end_date, next_start, next_end)
    week_id = week["id"]
    print(f"[OK] report_week #{week_id}: {start_date} ~ {end_date} "
          f"(차주 {next_start} ~ {next_end})")

    # ── 운영 현황 (rows 4-6; row 7 '총계' is a computed total, not imported) ──
    op_count = 0
    for row in (4, 5, 6):
        client = _s(ws.cell(row=row, column=2).value)
        if not client:
            continue
        report_store.add_operation_status(
            week_id, client, op_count,
            last_web=_n(ws.cell(row=row, column=3).value),
            last_was=_n(ws.cell(row=row, column=4).value),
            last_dev=_n(ws.cell(row=row, column=5).value),
            this_web=_n(ws.cell(row=row, column=7).value),
            this_was=_n(ws.cell(row=row, column=8).value),
            this_dev=_n(ws.cell(row=row, column=9).value),
        )
        op_count += 1
    print(f"[OK] operation_status: {op_count}건")

    # ── 금주 진행 내역 (header row 10, data starts row 11) ──────────────────
    cur_count = 0
    row = 11
    while _s(ws.cell(row=row, column=2).value):
        report_store.add_work_log(
            week_id, "current", cur_count,
            category=_s(ws.cell(row=row, column=3).value),
            request_date=_s_date(ws.cell(row=row, column=4).value),
            requester=_s(ws.cell(row=row, column=5).value),
            work_date=_s_date(ws.cell(row=row, column=6).value),
            work_content=_s(ws.cell(row=row, column=7).value),
            web_count=_n(ws.cell(row=row, column=11).value),
            was_count=_n(ws.cell(row=row, column=12).value),
            etc_count=_n(ws.cell(row=row, column=13).value),
        )
        cur_count += 1
        row += 1
    print(f"[OK] work_log (금주): {cur_count}건")

    # ── 차주 진행 내역 (header row 28, data starts row 29) ──────────────────
    # 차주 표는 WEB/WAS 열이 없고 "기타" 하나만 있으며, 그 열이 K(11)에 위치한다.
    next_count = 0
    row = 29
    while _s(ws.cell(row=row, column=2).value):
        report_store.add_work_log(
            week_id, "next", next_count,
            category=_s(ws.cell(row=row, column=3).value),
            request_date=_s_date(ws.cell(row=row, column=4).value),
            requester=_s(ws.cell(row=row, column=5).value),
            work_date=_s_date(ws.cell(row=row, column=6).value),
            work_content=_s(ws.cell(row=row, column=7).value),
            web_count=0,
            was_count=0,
            etc_count=_n(ws.cell(row=row, column=11).value),
        )
        next_count += 1
        row += 1
    print(f"[OK] work_log (차주): {next_count}건")

    # ── 특이 사항 (header row 32, data starts row 33) ───────────────────────
    note_count = 0
    row = 33
    while _s(ws.cell(row=row, column=2).value):
        report_store.add_special_note(
            week_id, note_count,
            category=_s(ws.cell(row=row, column=3).value),
            note_date=_s_date(ws.cell(row=row, column=4).value),
            title=_s(ws.cell(row=row, column=5).value),
            service_name=_s(ws.cell(row=row, column=6).value),
            detail=_s(ws.cell(row=row, column=7).value),
        )
        note_count += 1
        row += 1
    print(f"[OK] special_note: {note_count}건")

    return 0


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("Usage: python3 seed_report_from_xlsx.py /path/to/report.xlsx")
        sys.exit(1)
    report_store.init_db()
    sys.exit(seed(sys.argv[1]))
